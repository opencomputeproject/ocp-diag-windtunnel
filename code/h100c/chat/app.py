# app.py — single-user local chatbot wrapping the H100 vLLM endpoint.
#
# OpenAI-compatible backend (Qwen/Gemma/… on :8000) + SQLite-backed threaded
# history: create / list / open / rename / delete conversations, streamed replies.
# Single user, no auth — meant to sit local to the box (or behind a tunnel).

import json
import io
import os
import re
import sqlite3
import time
import uuid
from contextlib import closing

import httpx
from fastapi import FastAPI, HTTPException, Query
from fastapi.responses import FileResponse, Response, StreamingResponse
from pydantic import BaseModel

VLLM_BASE = os.environ.get("VLLM_BASE", "http://localhost:8000/v1").rstrip("/")
VLLM_API_KEY = os.environ.get("VLLM_API_KEY", "h100")
DB_PATH = os.environ.get("CHAT_DB", "/data/db/chat.db")
SYSTEM_PROMPT = os.environ.get("CHAT_SYSTEM", "").strip()
HERE = os.path.dirname(os.path.abspath(__file__))

app = FastAPI(title="H100 Chat")


# --- storage ---------------------------------------------------------------

def db():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA foreign_keys = ON")
    return conn


def init_db():
    os.makedirs(os.path.dirname(DB_PATH) or ".", exist_ok=True)
    with closing(db()) as conn, conn:
        conn.execute(
            """CREATE TABLE IF NOT EXISTS threads (
                   id TEXT PRIMARY KEY,
                   title TEXT NOT NULL,
                   created REAL NOT NULL,
                   updated REAL NOT NULL)"""
        )
        conn.execute(
            """CREATE TABLE IF NOT EXISTS messages (
                   seq INTEGER PRIMARY KEY AUTOINCREMENT,
                   thread_id TEXT NOT NULL,
                   role TEXT NOT NULL,
                   content TEXT NOT NULL,
                   created REAL NOT NULL,
                   FOREIGN KEY (thread_id) REFERENCES threads(id) ON DELETE CASCADE)"""
        )
        conn.execute(
            "CREATE INDEX IF NOT EXISTS idx_msg_thread ON messages(thread_id, seq)"
        )
        conn.execute(
            """CREATE TABLE IF NOT EXISTS artifacts (
                   id INTEGER PRIMARY KEY AUTOINCREMENT,
                   thread_id TEXT NOT NULL,
                   seq INTEGER NOT NULL,
                   format TEXT NOT NULL,
                   label TEXT NOT NULL,
                   title TEXT NOT NULL,
                   content TEXT NOT NULL,
                   created REAL NOT NULL,
                   FOREIGN KEY (thread_id) REFERENCES threads(id) ON DELETE CASCADE)"""
        )
        conn.execute(
            "CREATE INDEX IF NOT EXISTS idx_art_thread ON artifacts(thread_id, seq)"
        )


init_db()


# --- models / schemas ------------------------------------------------------

class NewThread(BaseModel):
    title: str | None = None


class RenameThread(BaseModel):
    title: str


class ChatIn(BaseModel):
    content: str
    model: str
    hide_thinking: bool = True


def row_to_thread(r):
    return {"id": r["id"], "title": r["title"], "created": r["created"], "updated": r["updated"]}


# --- thread CRUD -----------------------------------------------------------

@app.get("/api/models")
async def list_models():
    try:
        async with httpx.AsyncClient(timeout=10) as client:
            r = await client.get(
                f"{VLLM_BASE}/models",
                headers={"Authorization": f"Bearer {VLLM_API_KEY}"},
            )
            r.raise_for_status()
            ids = [m["id"] for m in r.json().get("data", [])]
            return {"models": ids, "endpoint": VLLM_BASE}
    except Exception as e:
        return {"models": [], "endpoint": VLLM_BASE, "error": str(e)}


@app.get("/api/threads")
def get_threads():
    with closing(db()) as conn:
        rows = conn.execute(
            "SELECT * FROM threads ORDER BY updated DESC"
        ).fetchall()
        return [row_to_thread(r) for r in rows]


@app.post("/api/threads")
def create_thread(body: NewThread):
    tid = uuid.uuid4().hex
    now = time.time()
    title = (body.title or "New chat").strip() or "New chat"
    with closing(db()) as conn, conn:
        conn.execute(
            "INSERT INTO threads (id, title, created, updated) VALUES (?,?,?,?)",
            (tid, title, now, now),
        )
    return {"id": tid, "title": title, "created": now, "updated": now}


@app.get("/api/threads/{tid}")
def get_thread(tid: str):
    with closing(db()) as conn:
        t = conn.execute("SELECT * FROM threads WHERE id=?", (tid,)).fetchone()
        if not t:
            raise HTTPException(404, "thread not found")
        msgs = conn.execute(
            "SELECT seq, role, content, created FROM messages WHERE thread_id=? ORDER BY seq",
            (tid,),
        ).fetchall()
        arts = conn.execute(
            "SELECT id, seq, format, label, title FROM artifacts WHERE thread_id=? ORDER BY id",
            (tid,),
        ).fetchall()
        by_seq = {}
        for a in arts:
            by_seq.setdefault(a["seq"], []).append({
                "format": a["format"], "label": a["label"], "title": a["title"],
                "url": _artifact_url(a["id"], a["format"]),
            })
        return {
            "thread": row_to_thread(t),
            "messages": [
                {**dict(m), "artifacts": by_seq.get(m["seq"], [])} for m in msgs
            ],
        }


@app.patch("/api/threads/{tid}")
def rename_thread(tid: str, body: RenameThread):
    title = body.title.strip() or "New chat"
    with closing(db()) as conn, conn:
        cur = conn.execute(
            "UPDATE threads SET title=?, updated=? WHERE id=?",
            (title, time.time(), tid),
        )
        if cur.rowcount == 0:
            raise HTTPException(404, "thread not found")
    return {"ok": True, "title": title}


@app.delete("/api/threads/{tid}")
def delete_thread(tid: str):
    with closing(db()) as conn, conn:
        cur = conn.execute("DELETE FROM threads WHERE id=?", (tid,))
        if cur.rowcount == 0:
            raise HTTPException(404, "thread not found")
    return {"ok": True}


# --- chat (streaming proxy) ------------------------------------------------

def _history(conn, tid):
    rows = conn.execute(
        "SELECT role, content FROM messages WHERE thread_id=? ORDER BY seq",
        (tid,),
    ).fetchall()
    return [{"role": r["role"], "content": r["content"]} for r in rows]


@app.post("/api/threads/{tid}/chat")
async def chat(tid: str, body: ChatIn):
    now = time.time()
    with closing(db()) as conn, conn:
        t = conn.execute("SELECT * FROM threads WHERE id=?", (tid,)).fetchone()
        if not t:
            raise HTTPException(404, "thread not found")
        conn.execute(
            "INSERT INTO messages (thread_id, role, content, created) VALUES (?,?,?,?)",
            (tid, "user", body.content, now),
        )
        # First user turn becomes the thread title until renamed.
        first = conn.execute(
            "SELECT COUNT(*) c FROM messages WHERE thread_id=? AND role='user'",
            (tid,),
        ).fetchone()["c"]
        if first == 1 and (t["title"] == "New chat"):
            auto = body.content.strip().splitlines()[0][:48] or "New chat"
            conn.execute("UPDATE threads SET title=? WHERE id=?", (auto, tid))
        history = _history(conn, tid)

    want = _detect_artifact(body.content)
    if want:
        gen = _artifact_stream(tid, want, body.model, history, body.content)
    else:
        gen = _chat_stream(tid, body, history)
    return StreamingResponse(gen, media_type="text/event-stream")


async def _chat_stream(tid, body, history):
    messages = ([{"role": "system", "content": SYSTEM_PROMPT}] if SYSTEM_PROMPT else []) + history
    payload = {"model": body.model, "messages": messages, "stream": True}
    if body.hide_thinking and "qwen" in body.model.lower():
        payload["chat_template_kwargs"] = {"enable_thinking": False}

    acc = []
    try:
        async with httpx.AsyncClient(timeout=None) as client:
            async with client.stream(
                "POST", f"{VLLM_BASE}/chat/completions",
                headers={"Authorization": f"Bearer {VLLM_API_KEY}"}, json=payload,
            ) as resp:
                if resp.status_code != 200:
                    detail = (await resp.aread()).decode("utf-8", "replace")
                    yield _sse({"error": f"{resp.status_code}: {detail[:500]}"})
                    return
                async for line in resp.aiter_lines():
                    if not line.startswith("data:"):
                        continue
                    data = line[5:].strip()
                    if data == "[DONE]":
                        break
                    try:
                        delta = json.loads(data)["choices"][0]["delta"]
                    except (json.JSONDecodeError, KeyError, IndexError):
                        continue
                    piece = delta.get("content")
                    if piece:
                        acc.append(piece)
                        yield _sse({"delta": piece})
    except Exception as e:
        yield _sse({"error": str(e)})
    finally:
        text = "".join(acc)
        if text:
            with closing(db()) as conn, conn:
                conn.execute(
                    "INSERT INTO messages (thread_id, role, content, created) VALUES (?,?,?,?)",
                    (tid, "assistant", text, time.time()),
                )
                conn.execute("UPDATE threads SET updated=? WHERE id=?", (time.time(), tid))
        yield _sse({"done": True})


async def _artifact_stream(tid, want, model, history, request):
    fmt, label = want["format"], want["label"]
    yield _sse({"status": f"Designing your {label}…"})
    try:
        spec = await _generate_spec(fmt, model, history, request)
    except Exception as e:
        yield _sse({"error": f"Couldn't build the {label}: {e}"})
        yield _sse({"done": True})
        return

    title = (spec.get("title") or request).strip()[:120] or "Document"
    units = spec.get("slides") if fmt == "pptx" else spec.get("sections")
    n = len(units or [])
    noun = "slide" if fmt == "pptx" else "section"
    summary = f'Done — "{title}" ({n} {noun}{"s" if n != 1 else ""}). Download it below.'

    with closing(db()) as conn, conn:
        cur = conn.execute(
            "INSERT INTO messages (thread_id, role, content, created) VALUES (?,?,?,?)",
            (tid, "assistant", summary, time.time()),
        )
        seq = cur.lastrowid
        conn.execute("UPDATE threads SET updated=? WHERE id=?", (time.time(), tid))
        ca = conn.execute(
            "INSERT INTO artifacts (thread_id, seq, format, label, title, content, created)"
            " VALUES (?,?,?,?,?,?,?)",
            (tid, seq, fmt, label, title, json.dumps(spec), time.time()),
        )
        aid = ca.lastrowid

    yield _sse({"delta": summary})
    yield _sse({"artifact": {"format": fmt, "label": label, "url": _artifact_url(aid, fmt)}})
    yield _sse({"done": True})


def _sse(obj):
    return f"data: {json.dumps(obj)}\n\n"


# --- artifact generation (model -> structured spec) ------------------------

_PPTX_PROMPT = (
    "You are a presentation designer. Using the conversation above as source material, create a "
    "professional slide deck for the user's request below.\n"
    'Output ONLY a JSON object (no prose, no code fences) of exactly this shape:\n'
    '{"title": str, "subtitle": str, '
    '"slides": [{"title": str, "bullets": [str, ...], "notes": str}]}\n'
    "Rules: 6-9 slides. Open with an overview/agenda slide and close with a summary or next-steps "
    "slide. 3-6 bullets per slide, each a short phrase (<= ~12 words), not full sentences. "
    "'notes' is 1-2 sentences of speaker notes. Be substantive and specific — do not restate the "
    "chat verbatim. JSON only.\nRequest: <REQUEST>"
)
_DOC_PROMPT = (
    "You are a professional writer. Using the conversation above as source material, write a "
    "polished document for the user's request below.\n"
    'Output ONLY a JSON object (no prose, no code fences) of exactly this shape:\n'
    '{"title": str, "subtitle": str, '
    '"sections": [{"heading": str, "paragraphs": [str, ...], "bullets": [str, ...]}]}\n'
    "Rules: start with a brief executive-summary section, then 3-6 well-developed sections with "
    "clear headings. Use 'paragraphs' for prose and 'bullets' for lists (either may be empty). "
    "Professional, self-contained — do not reference 'the chat'. JSON only.\nRequest: <REQUEST>"
)


async def _complete(model, messages):
    payload = {"model": model, "messages": messages, "stream": False, "temperature": 0.4}
    if "qwen" in model.lower():
        payload["chat_template_kwargs"] = {"enable_thinking": False}
    async with httpx.AsyncClient(timeout=None) as client:
        r = await client.post(
            f"{VLLM_BASE}/chat/completions",
            headers={"Authorization": f"Bearer {VLLM_API_KEY}"}, json=payload,
        )
        r.raise_for_status()
        return r.json()["choices"][0]["message"]["content"]


def _parse_spec(raw, fmt, request):
    m = re.search(r"\{.*\}", raw, re.S)
    try:
        spec = json.loads(m.group(0) if m else raw)
        if isinstance(spec, dict) and (spec.get("slides") or spec.get("sections")):
            return spec
    except (json.JSONDecodeError, AttributeError):
        pass
    # Fallback: wrap whatever text came back into a minimal valid spec.
    lines = [l.strip(" -*") for l in raw.splitlines() if l.strip()]
    if fmt == "pptx":
        return {"title": request[:80], "subtitle": "",
                "slides": [{"title": "Summary", "bullets": lines[:6], "notes": ""}]}
    return {"title": request[:80], "subtitle": "",
            "sections": [{"heading": "Summary", "paragraphs": lines[:8], "bullets": []}]}


async def _generate_spec(fmt, model, history, request):
    prompt = (_PPTX_PROMPT if fmt == "pptx" else _DOC_PROMPT).replace("<REQUEST>", request)
    raw = await _complete(model, history + [{"role": "user", "content": prompt}])
    return _parse_spec(raw, fmt, request)


def _artifact_url(aid, fmt):
    return f"/api/artifacts/{aid}.{fmt}"


# --- export / rendering ----------------------------------------------------

_FMT_LABEL = {"pdf": "PDF", "pptx": "PowerPoint deck", "md": "Markdown file"}
_MIME = {
    "md": "text/markdown",
    "pdf": "application/pdf",
    "pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
}

# Intent detection: "make this a pdf", "turn it into a slide deck", etc.
_ACTION = re.compile(
    r"\b(make|create|generate|export|convert|turn|save|download|produce|prepare|"
    r"build|draft|render|give me|i want|can i get|as an?|into an?|in an?|to an?)\b",
    re.I,
)
_FMT_PATTERNS = [
    ("pptx", re.compile(r"\b(pptx?|powerpoint|slides?|slide\s*deck|deck|presentation)\b", re.I)),
    ("pdf", re.compile(r"\bpdfs?\b", re.I)),
    ("md", re.compile(r"\b(markdown|\.md\b|md\s*file)\b", re.I)),
]


def _detect_artifact(text):
    if not _ACTION.search(text):
        return None
    for fmt, pat in _FMT_PATTERNS:
        if pat.search(text):
            return {"format": fmt, "label": _FMT_LABEL[fmt]}
    return None


def _stem(title):
    return re.sub(r"[^A-Za-z0-9_-]+", "_", title).strip("_") or "conversation"


def _attachment(data, ext, stem):
    return Response(
        content=data, media_type=_MIME[ext],
        headers={"Content-Disposition": f'attachment; filename="{stem}.{ext}"'},
    )


def _render(ext, title, md):
    ext = ext.lower()
    if ext == "md":
        return _attachment(md.encode("utf-8"), ext, _stem(title))
    if ext == "pdf":
        return _attachment(_md_to_pdf(md), ext, _stem(title))
    if ext == "pptx":
        return _attachment(_md_to_pptx(title, md), ext, _stem(title))
    raise HTTPException(400, "unknown format (use md, pdf, or pptx)")


def _transcript_md(title, rows, scope):
    lines = [f"# {title}", ""]
    for r in rows:
        if scope == "assistant" and r["role"] != "assistant":
            continue
        if scope == "assistant":
            lines += [r["content"], ""]
        else:
            lines += [f"## {'You' if r['role'] == 'user' else 'Assistant'}", "", r["content"], ""]
    return "\n".join(lines)


def _md_to_pdf(md_text):
    import markdown as mdlib
    from xhtml2pdf import pisa

    body = mdlib.markdown(md_text, extensions=["fenced_code", "tables"])
    html = (
        "<html><head><style>"
        "body{font-family:Helvetica;font-size:11pt;line-height:1.4;}"
        "h1{font-size:18pt;} h2{font-size:13pt;color:#333;margin-top:14pt;}"
        "pre,code{font-family:Courier;background:#f4f4f4;font-size:9pt;}"
        "pre{padding:6pt;} table{border-collapse:collapse;} "
        "td,th{border:1px solid #ccc;padding:3pt;}"
        f"</style></head><body>{body}</body></html>"
    )
    buf = io.BytesIO()
    pisa.CreatePDF(src=html, dest=buf)
    return buf.getvalue()


def _sections(md_text):
    """Split markdown into (heading, [bullet lines]) chunks, <=10 bullets each."""
    secs, head, body, in_code = [], None, [], False
    for raw in md_text.splitlines():
        line = raw.rstrip()
        if line.strip().startswith("```"):
            in_code = not in_code
            continue
        if in_code or not line.strip():
            continue
        if line.startswith("#"):
            if head is not None or body:
                secs.append((head or "", body))
            head, body = line.lstrip("#").strip(), []
        else:
            txt = re.sub(r"^[-*>\s]+", "", line).strip()
            if txt:
                body.append(txt)
    if head is not None or body:
        secs.append((head or "", body))
    out = []
    for h, b in secs:
        if not h and not b:
            continue
        for i in range(0, max(len(b), 1), 10):
            out.append((h if i == 0 else f"{h} (cont.)", b[i:i + 10]))
    return out


def _md_to_pptx(title, md_text):
    from pptx import Presentation

    prs = Presentation()
    cover = prs.slides.add_slide(prs.slide_layouts[0])
    cover.shapes.title.text = title[:120] or "Conversation"
    for head, bullets in _sections(md_text):
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = (head or title)[:120]
        tf = slide.placeholders[1].text_frame
        tf.clear()
        for i, b in enumerate(bullets):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = b[:220]
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


# --- themed renderers (structured spec -> polished file) -------------------

def _spec_to_md(spec):
    out = [f"# {spec.get('title', 'Untitled')}", ""]
    if spec.get("subtitle"):
        out += [f"*{spec['subtitle']}*", ""]
    for sec in spec.get("sections", []):
        out += [f"## {sec.get('heading', '')}", ""]
        for para in sec.get("paragraphs", []):
            out += [para, ""]
        for b in sec.get("bullets", []):
            out.append(f"- {b}")
        if sec.get("bullets"):
            out.append("")
    return "\n".join(out).rstrip() + "\n"


def _spec_to_pdf(spec):
    import html as htmlmod
    from xhtml2pdf import pisa

    esc = htmlmod.escape
    p = [
        "<html><head><style>",
        "@page{margin:2cm;}",
        "body{font-family:Helvetica;color:#1a1f2b;font-size:11pt;line-height:1.5;}",
        "h1.cover{font-size:30pt;color:#0f172a;margin-top:210pt;text-align:center;}",
        ".sub{font-size:14pt;color:#6b7280;text-align:center;margin-top:10pt;}",
        "h2{font-size:15pt;color:#2f5bd7;border-bottom:1.5px solid #2f5bd7;"
        "padding-bottom:3pt;margin-top:20pt;}",
        "p{text-align:justify;margin:6pt 0;}",
        "ul{margin:6pt 0 6pt 14pt;} li{margin:3pt 0;}",
        ".brk{page-break-after:always;}",
        "</style></head><body>",
        f'<div class="brk"><h1 class="cover">{esc(spec.get("title", "Untitled"))}</h1>',
    ]
    if spec.get("subtitle"):
        p.append(f'<div class="sub">{esc(spec["subtitle"])}</div>')
    p.append("</div>")
    for sec in spec.get("sections", []):
        p.append(f'<h2>{esc(sec.get("heading", ""))}</h2>')
        for para in sec.get("paragraphs", []):
            p.append(f"<p>{esc(para)}</p>")
        if sec.get("bullets"):
            p.append("<ul>")
            p += [f"<li>{esc(b)}</li>" for b in sec["bullets"]]
            p.append("</ul>")
    p.append("</body></html>")
    buf = io.BytesIO()
    pisa.CreatePDF(src="".join(p), dest=buf)
    return buf.getvalue()


def _spec_to_pptx(spec):
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR
    from pptx.util import Inches, Pt

    ACCENT = RGBColor(0x2F, 0x5B, 0xD7)
    DARK = RGBColor(0x0F, 0x17, 0x2A)
    MUTED = RGBColor(0x6B, 0x72, 0x80)
    INK = RGBColor(0x1A, 0x1F, 0x2B)
    WHITE = RGBColor(0xFF, 0xFF, 0xFF)
    SUBTLE = RGBColor(0xAD, 0xB5, 0xC2)

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    SW, SH = prs.slide_width, prs.slide_height
    blank = prs.slide_layouts[6]

    def box(slide, l, t, w, h):
        tf = slide.shapes.add_textbox(l, t, w, h).text_frame
        tf.word_wrap = True
        return tf

    # Title slide — dark background with accent strip.
    s = prs.slides.add_slide(blank)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = DARK
    strip = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.9), Inches(2.4),
                               Inches(2.2), Inches(0.12))
    strip.fill.solid()
    strip.fill.fore_color.rgb = ACCENT
    strip.line.fill.background()
    tf = box(s, Inches(0.9), Inches(2.7), Inches(11.5), Inches(2.5))
    r = tf.paragraphs[0].add_run()
    r.text = spec.get("title", "Untitled")
    r.font.size = Pt(40)
    r.font.bold = True
    r.font.color.rgb = WHITE
    if spec.get("subtitle"):
        pr = tf.add_paragraph().add_run()
        pr.text = spec["subtitle"]
        pr.font.size = Pt(20)
        pr.font.color.rgb = SUBTLE

    for i, sl in enumerate(spec.get("slides", []), 1):
        s = prs.slides.add_slide(blank)
        bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, Inches(1.15))
        bar.fill.solid()
        bar.fill.fore_color.rgb = ACCENT
        bar.line.fill.background()
        btf = bar.text_frame
        btf.margin_left = Inches(0.6)
        btf.vertical_anchor = MSO_ANCHOR.MIDDLE
        br = btf.paragraphs[0].add_run()
        br.text = sl.get("title", "")
        br.font.size = Pt(26)
        br.font.bold = True
        br.font.color.rgb = WHITE
        tf = box(s, Inches(0.75), Inches(1.5), Inches(11.8), Inches(5.3))
        first = True
        for b in sl.get("bullets", []):
            p = tf.paragraphs[0] if first else tf.add_paragraph()
            first = False
            run = p.add_run()
            run.text = "\u2022  " + str(b)
            run.font.size = Pt(18)
            run.font.color.rgb = INK
            p.space_after = Pt(10)
        pn = box(s, SW - Inches(1.1), SH - Inches(0.55), Inches(0.8), Inches(0.3))
        pr = pn.paragraphs[0].add_run()
        pr.text = str(i)
        pr.font.size = Pt(10)
        pr.font.color.rgb = MUTED
        if sl.get("notes"):
            s.notes_slide.notes_text_frame.text = sl["notes"]

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


@app.get("/api/artifacts/{aid}.{ext}")
def download_artifact(aid: int, ext: str):
    with closing(db()) as conn:
        a = conn.execute(
            "SELECT format, title, content FROM artifacts WHERE id=?", (aid,)
        ).fetchone()
    if not a:
        raise HTTPException(404, "artifact not found")
    ext = ext.lower()
    if ext != a["format"]:
        raise HTTPException(400, "format mismatch")
    try:
        spec = json.loads(a["content"])
    except (json.JSONDecodeError, TypeError):
        raise HTTPException(500, "artifact is corrupt")
    if ext == "md":
        data = _spec_to_md(spec).encode("utf-8")
    elif ext == "pdf":
        data = _spec_to_pdf(spec)
    elif ext == "pptx":
        data = _spec_to_pptx(spec)
    else:
        raise HTTPException(400, "unknown format (use md, pdf, or pptx)")
    return _attachment(data, ext, _stem(a["title"] or "document"))


@app.get("/api/threads/{tid}/export.{ext}")
def export_thread(tid: str, ext: str, scope: str = Query("all")):
    with closing(db()) as conn:
        t = conn.execute("SELECT * FROM threads WHERE id=?", (tid,)).fetchone()
        if not t:
            raise HTTPException(404, "thread not found")
        rows = conn.execute(
            "SELECT role, content FROM messages WHERE thread_id=? ORDER BY seq", (tid,)
        ).fetchall()
    md = _transcript_md(t["title"], rows, scope)
    return _render(ext, t["title"], md)


# --- static UI -------------------------------------------------------------

@app.get("/")
def index():
    return FileResponse(os.path.join(HERE, "index.html"))
